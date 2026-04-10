#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Verify extraction results."""

import openpyxl

wb = openpyxl.load_workbook('black_tower_scenes_complete.xlsx')
ws = wb.active

with open('extraction_verification.txt', 'w', encoding='utf-8') as f:
    f.write("="*70 + "\n")
    f.write("EXTRACTION VERIFICATION REPORT\n")
    f.write("="*70 + "\n\n")

    f.write(f"Total rows: {ws.max_row} (including header)\n")
    f.write(f"Total scenes: {ws.max_row - 1}\n")
    f.write(f"Columns: {[cell.value for cell in ws[1]]}\n\n")

    # Find scene 279
    f.write("="*70 + "\n")
    f.write("SCENE 279 VERIFICATION:\n")
    f.write("="*70 + "\n")
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if row[0] == '279':
            f.write(f"ID: {row[0]}\n")
            f.write(f"Text: {row[1]}\n")
            f.write(f"Exits: {row[2]}\n")
            f.write(f"Items: {row[3]}\n")
            break

    f.write("\n" + "="*70 + "\n")
    f.write("SAMPLE SCENES WITH EXITS:\n")
    f.write("="*70 + "\n")

    # Find scenes with exits
    count = 0
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if row[2] and count < 10:  # Has exits
            f.write(f"\nScene {row[0]}:\n")
            f.write(f"  Exits: {row[2]}\n")
            f.write(f"  Text preview: {row[1][:100]}...\n")
            count += 1

    f.write("\n" + "="*70 + "\n")
    f.write("SAMPLE SCENES WITH ITEMS:\n")
    f.write("="*70 + "\n")

    # Find scenes with items
    count = 0
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if row[3] and count < 10:  # Has items
            f.write(f"\nScene {row[0]}:\n")
            f.write(f"  Items: {row[3]}\n")
            f.write(f"  Text preview: {row[1][:100]}...\n")
            count += 1

print("Verification report saved to: extraction_verification.txt")
