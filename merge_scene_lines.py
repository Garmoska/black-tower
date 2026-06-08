"""Merge multi-line scenes in black-tower.docx into single-paragraph scenes.

Each scene begins with a paragraph containing only a positive integer (the scene number).
All subsequent paragraphs until the next scene number are joined into one paragraph,
with line breaks replaced by spaces.
"""
import re
from pathlib import Path
from docx import Document

SOURCE = Path(r"C:\Leonov\black-tower\black-tower.docx")
DEST = Path(r"C:\Leonov\black-tower\black-tower-monolith.docx")

SCENE_HEADER = re.compile(r"^\s*\d+\s*$")


def main():
    src = Document(str(SOURCE))
    dst = Document()

    current_number = None
    current_lines = []

    def flush():
        if current_number is None:
            return
        dst.add_paragraph(current_number)
        merged = " ".join(s.strip() for s in current_lines if s.strip())
        merged = re.sub(r"\s+", " ", merged).strip()
        dst.add_paragraph(merged)

    pre_scene_lines = []
    for para in src.paragraphs:
        text = para.text
        if SCENE_HEADER.match(text):
            if current_number is None:
                # Flush any pre-scene preamble as-is
                for line in pre_scene_lines:
                    dst.add_paragraph(line)
                pre_scene_lines = []
            else:
                flush()
            current_number = text.strip()
            current_lines = []
        else:
            if current_number is None:
                pre_scene_lines.append(text)
            else:
                current_lines.append(text)

    flush()
    dst.save(str(DEST))
    print(f"Saved: {DEST}")


if __name__ == "__main__":
    main()
